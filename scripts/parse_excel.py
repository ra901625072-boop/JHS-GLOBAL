import openpyxl

excel_path = r'client info and reqvierments\Products_website.xlsx'
wb = openpyxl.load_workbook(excel_path)
sheet = wb.active

print(f"Sheet name: {sheet.title}")
print(f"Max row: {sheet.max_row}, Max column: {sheet.max_column}")

headers = [sheet.cell(1, c).value for c in range(1, sheet.max_column + 1)]
print("Headers:")
for i, h in enumerate(headers, 1):
    print(f"  Col {i}: {repr(h)}")

print("\n--- Column Content Details ---")
for c in range(1, sheet.max_column + 1):
    header = sheet.cell(1, c).value
    rows = []
    for r in range(2, sheet.max_row + 1):
        val = sheet.cell(r, c).value
        if val is not None and str(val).strip() != "":
            rows.append((r, str(val).strip()))
    print(f"\nCol {c}: {repr(header)} (Count: {len(rows)})")
    for r_idx, val in rows:
        print(f"  Row {r_idx:2d}: {val}")
